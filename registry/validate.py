"""Read-only checker for the tracking contract (docs/TRACKING_CONTRACT.md).

Scans the real artifact folders, reports per entity how many artifacts could be identified (id in
contract format) and linked (every required link resolves), and lists what is missing. Validates
JSONL files against registry/schemas. Writes nothing except the optional --report file, which is
refused under C:\\Users\\pujan\\Box\\ and C:\\Users\\pujan\\Desktop\\Cryo-image-processing\\.

    python registry/validate.py                       # default roots (siblings of the cryoEV checkout)
    python registry/validate.py --jsonl some.jsonl --report C:/tmp/tracking_report.json
    python registry/validate.py --laptop <CryoAI-laptop> --cache <CryoAI> --processing <Cryo-image-processing>

Session 20261012 (lipo/VLDL) belongs to another project: its images, grids and batches are skipped
and only counted as `excluded`, never named.
"""
import argparse
import csv
import hashlib
import json
import os
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import ids  # noqa: E402
from schema_check import check_record, load_schemas  # noqa: E402

REPO = HERE.parent
PROTECTED = [Path(r"C:\Users\pujan\Box"), Path(r"C:\Users\pujan\Desktop\Cryo-image-processing")]
EXCLUDED_SESSIONS = {"20261012"}
ENTITIES = ["image", "session", "grid", "roboflow_image", "pull", "roboflow_export", "label_version",
            "review_package", "inference_check", "dataset", "job_bundle", "run", "model"]
N_EXAMPLES = 3


# ---------------------------------------------------------------- helpers
def find_sibling(name, start=REPO):
    """First `<ancestor>/<name>` folder above `start` (the checkout may be a nested worktree)."""
    for p in [start, *start.parents]:
        if (p / name).is_dir():
            return p / name
    return None


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def read_json(path):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def read_datajs(path):
    s = Path(path).read_text(encoding="utf-8")
    return json.loads(s[s.index("{"):s.rindex("}") + 1])


def read_jsonl(path):
    """[(line_no, obj | None, error | None)]"""
    out = []
    with open(path, encoding="utf-8") as fh:
        for n, line in enumerate(fh, 1):
            if not line.strip():
                continue
            try:
                out.append((n, json.loads(line), None))
            except json.JSONDecodeError as e:
                out.append((n, None, f"unparseable: {e.msg}"))
    return out


def under(path, root):
    a = os.path.normcase(os.path.abspath(str(path)))
    b = os.path.normcase(os.path.abspath(str(root)))
    return a == b or a.startswith(b.rstrip("\\/") + os.sep)


def check_report_path(path, protected=PROTECTED):
    """Raise ValueError if the report may not be written there."""
    p = Path(os.path.realpath(os.path.abspath(str(path))))
    for root in protected:
        if under(p, root) or under(os.path.abspath(str(path)), root):
            raise ValueError(f"refused: --report {path} is under protected root {root}")
    if p.exists():
        try:
            if read_json(p).get("schema") != "validate_report/1":
                raise ValueError(f"refused: {p} exists and is not a previous validate report")
        except (OSError, ValueError, AttributeError):
            raise ValueError(f"refused: {p} exists and is not a previous validate report")
    return p


class Tally:
    """Per-entity counts: found / identified / linked / excluded, plus missing reasons with examples."""

    def __init__(self):
        self.found = self.identified = self.linked = self.excluded = 0
        self.missing = Counter()
        self.examples = defaultdict(list)
        self.notes = []

    def add(self, ident, ok_id, problems):
        self.found += 1
        self.identified += bool(ok_id)
        if not ok_id:
            problems = ["id not in contract format"] + list(problems)
        self.linked += not problems
        for p in problems:
            self.missing[p] += 1
            if len(self.examples[p]) < N_EXAMPLES:
                self.examples[p].append(ident)

    def as_dict(self):
        return {"found": self.found, "identified": self.identified, "linked": self.linked,
                "excluded": self.excluded, "missing": dict(self.missing.most_common()),
                "examples": dict(self.examples), "notes": self.notes}


# ---------------------------------------------------------------- scanners
class Scan:
    def __init__(self, laptop, cache, processing, repo=REPO, schemas=None):
        self.laptop, self.cache, self.proc, self.repo = (Path(p) if p else None for p in (laptop, cache, processing, repo))
        self.schemas = schemas or load_schemas()
        self.t = {e: Tally() for e in ENTITIES}
        self.jsonl = {}
        self.excluded_keys = set()      # image keys of excluded-session images
        self.image_info = {}            # image key -> dict
        self.key_sessions = defaultdict(set)   # image key -> session folders whose MANIFESTs list it
        self.pulls, self.exports, self.packages, self.models, self.runs, self.datasets = {}, {}, {}, {}, {}, {}
        self.bundles, self.sessions, self.grids, self.slot_grid = {}, {}, {}, {}

    def ok(self, root):
        return root is not None and Path(root).is_dir()

    # -- order matters: exclusions and canonical stems first
    def run(self):
        self.load_exclusions()
        self.load_models()
        self.scan_pulls()
        self.scan_tracking()
        self.scan_master_and_tracker()
        self.scan_session_tiles()
        self.scan_exports()
        self.scan_runs_and_datasets()
        self.scan_bundles()
        self.scan_packages()
        self.scan_label_versions()
        self.scan_inference_checks()
        self.finish_images()
        self.finish_models()
        return self

    def is_excluded_stem(self, stem):
        return ids.image_key(stem) in self.excluded_keys

    def load_exclusions(self):
        if not self.ok(self.cache):
            return
        for f in (self.cache / "training images").glob("*/split_assignment.csv"):
            for r in read_csv(f):
                if (r.get("batch") or "")[:8] in EXCLUDED_SESSIONS:
                    self.excluded_keys.add(ids.image_key(r["image"]))

    def canon(self, stem):
        key = ids.image_key(stem)
        return key if key in self.image_info else None

    def note_stem(self, stem, **info):
        if self.is_excluded_stem(stem):
            return None
        canon = ids.image_key(stem)
        d = self.image_info.setdefault(canon, {})
        for k, v in info.items():
            if v not in (None, ""):
                d.setdefault(k, v)
        return canon

    def load_models(self):
        import yaml
        path = self.repo / "models" / "models.yaml"
        if path.exists():
            self.models = (yaml.safe_load(path.read_text(encoding="utf-8")) or {}).get("models", {}) or {}

    def scan_pulls(self):
        t = self.t["pull"]
        if not self.ok(self.laptop):
            t.notes.append("CryoAI-laptop not found")
            return
        for d in sorted(self.laptop.glob("roboflow_pull_*")):
            coco = d / "current.coco.json"
            probs = []
            stems = {}
            if coco.exists():
                c = read_json(coco)
                for im in c["images"]:
                    ex = im.get("extra", {})
                    stem = ids.image_key(ex.get("name") or im["file_name"])
                    if self.is_excluded_stem(stem):
                        t.excluded += 1
                        continue
                    stems[stem] = ex.get("roboflow_id")
                    self.note_stem(stem, rf_id=ex.get("roboflow_id"), rf_name=im["file_name"], pull=d.name,
                                   split=ex.get("split"))
                    rt = self.t["roboflow_image"]
                    rid = ex.get("roboflow_id")
                    rt.add(rid or stem, ids.is_id("rf_id", rid), [] if rid else ["no roboflow_id in pull"])
                if any(v is None for v in stems.values()):
                    probs.append("images without roboflow_id")
            else:
                probs.append("current.coco.json missing")
            if not (d / "roboflow_state.json").exists():
                probs.append("roboflow_state.json missing (project/created not recorded)")
            self.pulls[d.name] = {"path": d, "stems": stems, "coco": coco if coco.exists() else None}
            t.add(d.name, ids.is_id("pull", d.name), probs)
            t.notes.append(f"{d.name}: {len(stems)} images; project name is not stored in the pull folder")

    def scan_tracking(self):
        tr = self.laptop / "tracking" / "v2" if self.ok(self.laptop) else None
        if not tr or not (tr / "prelipo_tracking.csv").exists():
            self.t["image"].notes.append("tracking/v2/prelipo_tracking.csv not found")
            return
        for r in read_csv(tr / "prelipo_tracking.csv"):
            self.note_stem(r["stem"], rf_id=r.get("roboflow_id"), session_date=r.get("session_date"),
                           session_date_basis=(r.get("session_date_basis") or "").split(":")[0],
                           session_label=r.get("session"), grid=r.get("grid_id"), session_grid=r.get("grid"),
                           duplicate_of=r.get("duplicate_of"), tracking="prelipo_tracking.csv")

    def scan_master_and_tracker(self):
        st, gt = self.t["session"], self.t["grid"]
        master = (self.proc / "imaged_grids_master.csv") if self.ok(self.proc) else None
        sess = defaultdict(lambda: {"master": False, "tracker": False, "folder": False, "grids": set()})
        grids = defaultdict(lambda: {"master": False, "tracker": False, "sessions": set()})
        no_inv = Counter()
        if master and master.exists():
            for r in read_csv(master):
                sid = (r.get("session_date") or "").replace("-", "")
                if sid[:8] in EXCLUDED_SESSIONS:
                    continue
                sess[sid]["master"] = True
                inv = (r.get("inventory_id") or "").strip()
                if inv:
                    sess[sid]["grids"].add(inv)
                    grids[inv]["master"] = True
                    grids[inv]["sessions"].add(sid)
                    self.slot_grid[(sid, (r.get("session_grid") or "").strip())] = inv
                else:
                    no_inv[sid] += 1
        else:
            st.notes.append("imaged_grids_master.csv not found")
        xlsx = self.proc / "PKD_AU Cryo Tracker - 10-23-26 onwards.xlsx" if self.ok(self.proc) else None
        legacy = {}
        if xlsx and xlsx.exists():
            import openpyxl
            wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
            try:
                rows = list(wb["Cryo Grids"].iter_rows(min_row=3, values_only=True))
                for r in rows:
                    if isinstance(r[0], (int, float)):
                        g = str(int(r[0]))
                        img_date = r[15]
                        if isinstance(img_date, datetime) and img_date.strftime("%Y%m%d") in EXCLUDED_SESSIONS:
                            gt.excluded += 1
                            grids.pop(g, None)
                            continue
                        grids[g]["tracker"] = True
                for r in wb["Imaging Sessions"].iter_rows(min_row=2, values_only=True):
                    if isinstance(r[0], datetime):
                        sid = r[0].strftime("%Y%m%d")
                        if sid in EXCLUDED_SESSIONS:
                            continue
                        sess[sid]["tracker"] = True
                for r in wb["Grid ID Changes"].iter_rows(min_row=2, values_only=True):
                    if isinstance(r[0], (int, float)) and isinstance(r[1], (int, float)):
                        legacy[str(int(r[0]))] = str(int(r[1]))
            finally:
                wb.close()
        else:
            st.notes.append("tracker xlsx not found")
        if self.ok(self.proc):
            for d in self.proc.iterdir():
                if d.is_dir() and ids.is_id("session", d.name) and d.name[:8] not in EXCLUDED_SESSIONS:
                    sess[d.name]["folder"] = True
        for sid, s in sorted(sess.items()):
            probs = []
            base = sid[:8]
            in_tracker = s["tracker"] or sess.get(base, {}).get("tracker", False)
            if not in_tracker:
                probs.append("not in tracker 'Imaging Sessions'")
            if not s["master"] and not (base in sess and sess[base]["master"]):
                probs.append("not in imaged_grids_master.csv")
            if not s["folder"]:
                probs.append("no local session folder (pre-manifest or not processed)")
            if no_inv.get(sid):
                probs.append("session grid(s) without inventory_id in master")
            self.sessions[sid] = s
            st.add(sid, ids.is_id("session", sid), probs)
        for g, d in sorted(grids.items()):
            probs = []
            if not d["tracker"]:
                probs.append("inventory id not in tracker 'Cryo Grids'")
            if g in legacy:
                probs.append(f"renumbered (Grid ID Changes) -> needs replaced_by")
            self.grids[g] = d
            gt.add(g, ids.is_id("grid", g), probs)
        gt.notes.append(f"{sum(no_inv.values())} master rows have no inventory_id (scope slot only); "
                        f"{sum(1 for d in grids.values() if not d['master'])} inventory grids not in master (not imaged yet or pre-master)")

    def scan_session_tiles(self):
        if not self.ok(self.proc):
            return
        for sdir in sorted(self.proc.iterdir()):
            if not (sdir.is_dir() and ids.is_id("session", sdir.name)) or sdir.name[:8] in EXCLUDED_SESSIONS:
                continue
            ml = sdir / "5. Model + labeling"
            manifests = [mf for mf in (ml / "model_input_MANIFEST.csv", ml / "model_input" / "MANIFEST.csv")
                         if mf.exists()][:1]
            if (ml / "ROBOFLOW_UPLOADS.csv").exists():      # upload batches: their own MANIFEST with roboflow_id
                for u in read_csv(ml / "ROBOFLOW_UPLOADS.csv"):
                    m = Path(u.get("manifest") or "")
                    if m.name and m.exists() and under(m, self.proc) and m not in manifests:
                        manifests.append(m)
            for mf in manifests:
                upload = mf.parent.name != "5. Model + labeling" and mf.parent.name != "model_input"
                for r in read_csv(mf):
                    name = r.get("output_name") or r.get("tile") or ""
                    if not name:
                        continue
                    slot = (r.get("grid") or "").strip()
                    key = self.note_stem(name, session_date=sdir.name, session_date_basis="folder",
                                         session_grid=slot, grid=self.slot_grid.get((sdir.name, slot)),
                                         nm_per_px=r.get("nm_per_px"), md5=r.get("md5"), manifest=str(mf),
                                         rf_id=r.get("roboflow_id"), uploaded=upload or None)
                    if key:
                        self.key_sessions[key].add(sdir.name)
                        if upload and r.get("source_file") == name:
                            self.image_info[key]["copy_unlinked"] = True

    def scan_exports(self):
        t = self.t["roboflow_export"]
        if not self.ok(self.cache):
            t.notes.append("CryoAI cache not found")
            return
        for d in sorted((self.cache / "training images").iterdir()):
            if not d.is_dir():
                continue
            if d.name.startswith("_"):
                t.notes.append(f"skipped local working copy {d.name} (not an export)")
                continue
            counts, unmapped, probs = {}, 0, []
            for split in ("train", "valid", "test"):
                f = d / split / "_annotations.coco.json"
                if f.exists():
                    c = read_json(f)
                    keep = [im for im in c["images"] if not self.is_excluded_stem(ids.stem_of_export(im["file_name"]))]
                    t.excluded += len(c["images"]) - len(keep)
                    counts[split] = len(keep)
                    for im in keep:
                        if self.canon(ids.stem_of_export(im["file_name"])) is None:
                            unmapped += 1
            if not counts:
                probs.append("no local COCO (cache placeholder; export lives on Box/desktop)")
            if unmapped:
                probs.append("export images with no known canonical stem (not in a pull, tracking or MANIFEST)")
                t.notes.append(f"{d.name}: {unmapped} images not mapped to a canonical stem")
            derived = None
            rep = d / "SPLIT_REPORT.md"
            if rep.exists():
                m = re.search(r"Source: `([^`]+)`", rep.read_text(encoding="utf-8"))
                derived = m.group(1) if m else None
            self.exports[d.name] = {"counts": counts, "derived_from": derived,
                                    "split_file": d / "split_assignment.csv" if (d / "split_assignment.csv").exists() else None}
            t.add(d.name, ids.is_id("export", d.name), probs)

    def export_of(self, text):
        """Longest export id that prefixes a models.yaml `train_data` value."""
        hits = [e for e in self.exports if text and text.startswith(e)]
        if not hits and text:
            m = re.match(r"^(roboflow[^()]*?)\s*(\(|$)", text)
            return m.group(1).strip() if m else None
        return max(hits, key=len) if hits else None

    def scan_runs_and_datasets(self):
        rt, dt = self.t["run"], self.t["dataset"]
        refs = {}
        for mid, m in self.models.items():
            sr = (m or {}).get("source_run") or ""
            parts = Path(sr).parts
            if "runs" in parts:
                i = parts.index("runs")
                refs[parts[i + 1]] = {"dataset": parts[i - 1], "model": mid, "export": self.export_of(m.get("train_data"))}
        local = {}
        if self.ok(self.cache):
            for rd in (self.cache / "training outputs").glob("*/runs/*"):
                if rd.is_dir():
                    local[rd.name] = rd
        for rid in sorted(set(refs) | set(local)):
            ref = refs.get(rid, {})
            dataset = ref.get("dataset") or (local[rid].parent.parent.name if rid in local else None)
            probs = ["no job bundle (pre-bundle run)"]
            if rid not in local:
                probs.append("run folder not in local cache")
            elif not (local[rid] / "args.yaml").exists():
                probs.append("args.yaml/results.csv not in local cache (config from folder name only)")
            if not ids.parse_run(rid):
                probs.append("config not parseable from run name")
            if not ref.get("model"):
                probs.append("no registered model from this run")
            self.runs[rid] = {"dataset": dataset, "model": ref.get("model"), "local": rid in local}
            rt.add(rid, ids.is_id("run", rid), probs)
            ds = self.datasets.setdefault(dataset, {"runs": [], "export": None})
            ds["runs"].append(rid)
            ds["export"] = ds["export"] or ref.get("export")
        if self.ok(self.cache):
            for d in (self.cache / "training outputs").iterdir():
                if d.is_dir():
                    self.datasets.setdefault(d.name, {"runs": [], "export": None})
        for did, d in sorted(self.datasets.items()):
            probs = ["no label_version (trained from a Roboflow export)"]
            exp = d["export"]
            if not exp:
                probs.append("source export unknown")
            elif exp not in self.exports:
                probs.append("source export not in local cache")
            elif not self.exports[exp]["counts"]:
                probs.append("source export has no local COCO (image list not recoverable locally)")
            elif not self.exports[exp]["split_file"]:
                probs.append("split list = export split folders (no split file)")
            if not self.ok(self.cache) or not (self.cache / "training outputs" / did / "dataset.yaml").exists():
                probs.append("dataset.yaml not in local cache")
            dt.add(did, ids.is_id("dataset", did), probs)

    def scan_bundles(self):
        t = self.t["job_bundle"]
        d = self.laptop / "job_bundles" if self.ok(self.laptop) else None
        if not d or not d.is_dir():
            t.notes.append("CryoAI-laptop/job_bundles/ not found: no bundles made yet")
            return
        for z in sorted(d.glob("*.zip")):
            bid = z.stem
            probs = []
            try:
                with zipfile.ZipFile(z) as zf:
                    jobs = [n for n in zf.namelist() if n.endswith("/JOB.json")]
                    job = json.loads(zf.read(jobs[0])) if jobs else None
            except (zipfile.BadZipFile, OSError, json.JSONDecodeError):
                job = None
            if not job:
                probs.append("JOB.json missing or unreadable")
            elif ids.bundle_dataset(bid) != job.get("dataset_name"):
                probs.append("bundle id dataset != JOB.json dataset_name")
            ds = ids.bundle_dataset(bid)
            if ds and ds not in self.datasets:
                probs.append("dataset not known (no training outputs folder / run)")
            if not any(r["dataset"] == ds for r in self.runs.values()):
                probs.append("no run from this bundle yet")
            self.bundles[bid] = job
            t.add(bid, ids.is_id("bundle", bid), probs)

    def scan_packages(self):
        t = self.t["review_package"]
        found = []
        if self.ok(self.laptop):
            found += list((self.laptop / "review").glob("*/data.js")) + list((self.laptop / "review").glob("*/package/data.js"))
        if self.ok(self.cache):
            found += list((self.cache / "annotation review").glob("*/data.js"))
        by_id = defaultdict(list)
        metas = []
        scratch = [f for f in found if f.parent.name == "package" or f.parent.name.startswith("_")]
        if scratch:
            t.notes.append(f"skipped {len(scratch)} control/scratch package(s) (folders starting with _)")
        for f in sorted(set(found) - set(scratch)):
            folder = f.parent.name
            try:
                d = read_datajs(f)
            except (ValueError, OSError):
                t.add(folder, False, ["data.js unreadable"])
                continue
            metas.append((f, folder, d))
            by_id[d.get("id")].append(folder)
        run_to_model = {}
        for mid, m in self.models.items():
            run_to_model[Path((m or {}).get("source_run") or "x").name] = mid
        for f, folder, d in metas:
            pid = d.get("id")
            probs = []
            if len(by_id[pid]) > 1:
                probs.append("data.js id shared by several folders")
            if folder != pid:
                probs.append("folder name != data.js id")
            src = d.get("source") or ""
            m = re.search(r"(roboflow_pull_\d{8}_\d{4})", src)
            base = m.group(1) if m else src
            if m:
                if base not in self.pulls:
                    probs.append("base pull not found")
            elif base not in self.exports:
                probs.append("base export not found locally")
            run = (d.get("run") or "").split(" ")[0]
            if run not in self.models and run not in run_to_model:
                probs.append("model/run not resolvable")
            pstems = self.pulls.get(base, {}).get("stems", {})
            ed = f.parent / "edits"
            if ed.is_dir():
                recs = list(ed.glob("*.json"))
                legacy = bad_pkg = unknown = 0
                for r in recs:
                    try:
                        w = read_json(r)
                    except (ValueError, OSError):
                        legacy += 1
                        continue
                    if w.get("fmt") != 1:
                        legacy += 1
                    elif w.get("pkg") != pid:
                        bad_pkg += 1
                    if pstems and ids.image_key(w.get("file") or r.stem) not in pstems:
                        unknown += 1
                if legacy:
                    probs.append("edit records without pkg/fmt (pre-v1)")
                if bad_pkg:
                    probs.append("edit record pkg != data.js id")
                if unknown:
                    probs.append("edit record image not in base pull")
            self.packages.setdefault(pid, []).append({"folder": folder, "base": base, "path": f.parent})
            t.add(f"{folder} (id {pid})" if folder != pid else pid, ids.is_id("pkg", pid or ""), probs)
        t.notes.append("status (built/in_review/reviewed/applied) is not recorded anywhere yet; derived only")

    def scan_label_versions(self):
        t = self.t["label_version"]
        if not self.ok(self.laptop):
            return
        for d in sorted((self.laptop / "labels").glob("lv_*")):
            probs = []
            vj = d / "version.json"
            if not vj.exists():
                t.add(d.name, ids.is_id("lv", d.name), ["version.json missing"])
                continue
            v = read_json(vj)
            superseded = (d / "SUPERSEDED.md").exists()
            rec = {"schema": "label_version/1", "logged": v.get("created"),
                   "status": "superseded" if superseded else "created", **v}
            errs = check_record(rec, self.schemas)
            if errs:
                probs.append("version.json does not fit label_version/1")
                t.notes.append(f"{d.name}: " + "; ".join(errs[:3]))
            pull = self.pulls.get(v.get("base"))
            if not pull:
                probs.append("base pull not found")
            elif pull["coco"] and v.get("base_sha256") and sha256_file(pull["coco"]) != v["base_sha256"]:
                probs.append("base pull sha256 changed since the version was made")
            pk = self.packages.get(v.get("pkg"), [])
            if not pk:
                probs.append("pkg not found")
            elif len(pk) > 1:
                probs.append("pkg id ambiguous (several folders)")
            if v.get("id") != d.name:
                probs.append("version.json id != folder name")
            if superseded:
                t.notes.append(f"{d.name}: SUPERSEDED.md present")
            ch = d / "changes.jsonl"
            if ch.exists():
                self.validate_jsonl(ch)
                stems = (pull or {}).get("stems", {})
                bad = sum(1 for _, o, _ in read_jsonl(ch) if o and stems and o.get("image") not in stems)
                if bad:
                    probs.append("changes.jsonl rows whose image is not in the base pull")
            else:
                probs.append("changes.jsonl missing")
            t.add(d.name, ids.is_id("lv", d.name), probs)

    def scan_inference_checks(self):
        t = self.t["inference_check"]
        bj = self.laptop / "inference_checks" / "batches.jsonl" if self.ok(self.laptop) else None
        if not bj or not bj.exists():
            t.notes.append("inference_checks/batches.jsonl not found")
            return
        self.validate_jsonl(bj)
        last = {}
        for _, o, _ in read_jsonl(bj):
            if o and o.get("batch_id"):
                last[o["batch_id"]] = o
        for bid, o in sorted(last.items()):
            if (o.get("session") or "")[:8] in EXCLUDED_SESSIONS:
                t.excluded += 1
                continue
            probs = []
            if not (bj.parent / bid).is_dir():
                probs.append("batch folder missing")
            if o.get("model") not in self.models:
                probs.append("model not in models.yaml")
            if o.get("session") not in self.sessions:
                probs.append("session not known")
            if o.get("status") != "scored":
                t.notes.append(f"{bid}: status {o.get('status')}")
            t.add(bid, ids.is_id("ic", bid), probs)

    def finish_images(self):
        t, rt = self.t["image"], self.t["roboflow_image"]
        t.excluded = len(self.excluded_keys)
        basis = Counter()
        for stem, d in sorted(self.image_info.items()):
            probs = []
            b = d.get("session_date_basis") or "none"
            basis[b] += 1
            if not d.get("session_date"):
                probs.append("no session date")
            elif b == "inferred":
                probs.append("session date inferred, not from tracker")
            if not d.get("grid"):
                probs.append("no inventory Grid#")
            if d.get("uploaded") and not d.get("rf_id"):
                probs.append("uploaded session tile: Roboflow id unknown locally (needs a pull)")
            if len(self.key_sessions.get(stem, ())) > 1:
                probs.append("stem not unique across sessions")
            if d.get("copy_unlinked"):
                probs.append("upload copy (e.g. blind_NN): link to its source tile not in the MANIFEST")
            if d.get("pull") and not d.get("tracking"):
                probs.append("in pull but not in tracking backfill")
            t.add(stem, ids.is_id("stem", stem), probs)
        t.notes.append("session_date_basis: " + ", ".join(f"{k}={v}" for k, v in sorted(basis.items())))
        rt.notes.append("roboflow_image ids come from pulls only; upload batch/time are in prelipo_tracking.csv")

    def finish_models(self):
        t = self.t["model"]
        for mid, m in self.models.items():
            m = m or {}
            probs = []
            run = Path(m.get("source_run") or "").name or None
            if not run:
                probs.append("no source_run")
            elif run not in self.runs:
                probs.append("run not known")
            if not (self.repo / (m.get("weights") or "missing")).exists():
                probs.append("weights file not in this checkout")
            if not ids.is_id("model_new", mid):
                probs.append("legacy id (not v<N>_<tag>_<arch>_<YYYYMMDD>)")
            t.add(mid, ids.is_id("model", mid), probs)
        t.notes.append("models.yaml has no status / dataset id / weights sha256 fields")

    # -- JSONL
    def validate_jsonl(self, path):
        res = {"lines": 0, "valid": 0, "legacy_mapped": 0, "errors": [], "by_schema": Counter()}
        for n, o, err in read_jsonl(path):
            res["lines"] += 1
            if err:
                res["errors"].append(f"line {n}: {err}")
                continue
            if isinstance(o, dict) and (o.get("session") or "")[:8] in EXCLUDED_SESSIONS:
                res["lines"] -= 1
                res["excluded"] = res.get("excluded", 0) + 1
                continue
            if isinstance(o, dict) and "schema" not in o and "batch_id" in o:   # batches.jsonl v0 line
                o = {"schema": "inference_check/1", "id": o["batch_id"], **o}
                res["legacy_mapped"] += 1
            errs = check_record(o, self.schemas)
            res["by_schema"][o.get("schema") if isinstance(o, dict) else "?"] += 1
            if errs:
                res["errors"].append(f"line {n}: " + "; ".join(errs[:4]))
            else:
                res["valid"] += 1
        res["by_schema"] = dict(res["by_schema"])
        self.jsonl[str(path)] = res
        return res

    def summary(self):
        return {"schema": "validate_report/1", "created": datetime.now().isoformat(timespec="seconds"),
                "roots": {"laptop": str(self.laptop), "cache": str(self.cache), "processing": str(self.proc),
                          "repo": str(self.repo)},
                "excluded_sessions_policy": "skipped; counted only",
                "entities": {e: self.t[e].as_dict() for e in ENTITIES}, "jsonl": self.jsonl}


def print_summary(rep, out=None):
    w = (out or sys.stdout).write
    w(f"{'entity':16} {'found':>6} {'ident':>6} {'linked':>6} {'excl':>5}  top missing\n")
    for e, d in rep["entities"].items():
        top = "; ".join(f"{k} ({v})" for k, v in list(d["missing"].items())[:3])
        w(f"{e:16} {d['found']:>6} {d['identified']:>6} {d['linked']:>6} {d['excluded']:>5}  {top}\n")
    for p, r in rep["jsonl"].items():
        w(f"jsonl {Path(p).parent.name}/{Path(p).name}: {r['valid']}/{r['lines']} valid"
          f"{', legacy-mapped ' + str(r['legacy_mapped']) if r['legacy_mapped'] else ''}"
          f"{'; first error: ' + r['errors'][0] if r['errors'] else ''}\n")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--laptop", type=Path, default=find_sibling("CryoAI-laptop"))
    ap.add_argument("--cache", type=Path, default=find_sibling("CryoAI"))
    ap.add_argument("--processing", type=Path, default=Path(r"C:\Users\pujan\Desktop\Cryo-image-processing"))
    ap.add_argument("--repo", type=Path, default=REPO, help="cryoEV checkout holding models/models.yaml")
    ap.add_argument("--jsonl", type=Path, nargs="*", default=[], help="extra JSONL files to validate")
    ap.add_argument("--report", type=Path, help="write the full JSON report here (refused under protected roots)")
    ap.add_argument("--verbose", action="store_true", help="print examples and notes per entity")
    a = ap.parse_args(argv)
    report_path = check_report_path(a.report) if a.report else None   # refuse before any work
    s = Scan(a.laptop, a.cache, a.processing, a.repo).run()
    for p in a.jsonl:
        s.validate_jsonl(p)
    rep = s.summary()
    print_summary(rep)
    if a.verbose:
        for e, d in rep["entities"].items():
            for k, ex in d["examples"].items():
                print(f"  {e}: {k}: e.g. {', '.join(map(str, ex))}")
            for n in d["notes"]:
                print(f"  {e}: note: {n}")
    if report_path:
        report_path.write_text(json.dumps(rep, indent=1, default=str) + "\n", encoding="utf-8")
        print(f"report: {report_path}")
    bad = sum(len(r["errors"]) for r in rep["jsonl"].values())
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
