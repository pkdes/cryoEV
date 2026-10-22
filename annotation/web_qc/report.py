"""Human-readable QC outputs, built by annotation.web_qc.export.

  export/compare/<tile>.png   Original (model) | Revised, changes colour-coded
  export/compare_all.pdf      the same, one tile per page
  metrics/figures/*.png       P/R/F1 vs cut, precision by confidence bin,
                              class confusion, model-vs-reviewed EV sizes
  metrics/qc_metrics.xlsx     formatted tables
  metrics/qc_report.html      one self-contained page (figures + tables + thumbnails)

"Edited" is decided here, not stored by the tool: a kept model outline counts as
edited when its IoU with the model's original polygon is below EDIT_IOU. That
stays right after undo, or after an edit dragged back into place.
"""

from __future__ import annotations

import base64
import html
import io
import time
from collections import Counter
from pathlib import Path

import cv2
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
from PIL import Image, ImageDraw, ImageFont  # noqa: E402

from annotation.web_qc.server import read_predictions  # noqa: E402

EDIT_IOU = 0.98
CHANGE_RGB = {"reclassed": (255, 214, 10), "edited": (255, 140, 0), "drawn": (0, 220, 255), "rejected": (248, 81, 73),
              "pending": (139, 148, 158)}


def _rgb(hex_color):
    h = hex_color.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def _iou(a, b, w, h):
    ma, mb = np.zeros((h, w), np.uint8), np.zeros((h, w), np.uint8)
    cv2.fillPoly(ma, [np.round(a).astype(np.int32)], 1)
    cv2.fillPoly(mb, [np.round(b).astype(np.int32)], 1)
    u = np.logical_or(ma, mb).sum()
    return np.logical_and(ma, mb).sum() / u if u else 1.0


def classify_changes(d, pred_dir):
    """[(object, change)] where change is kept / reclassed / edited / drawn / rejected / pending."""
    w, h = d["width"], d["height"]
    orig, _ = read_predictions(pred_dir / (Path(d["image"]).stem + ".txt"), w, h)
    out = []
    for o in d["objects"]:
        if o["status"] == "pending":
            ch = "pending"
        elif o["status"] == "rejected":
            ch = "rejected"
        elif o["source"] == "human":
            ch = "drawn"
        else:
            k = int(o["id"][1:])
            cur = np.array([[x * w, y * h] for x, y in o["poly"]])
            if k < len(orig) and _iou(cur, orig[k], w, h) < EDIT_IOU:
                ch = "edited"
            elif o["cls"] != o["model_class"]:
                ch = "reclassed"
            else:
                ch = "kept"
        out.append((o, ch))
    return out


def _font(size):
    for f in ("segoeui.ttf", "arial.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(f, size)
        except OSError:
            continue
    return ImageFont.load_default()


def _poly(draw, pts, rgb, width, dashed=False):
    pts = [tuple(p) for p in pts] + [tuple(pts[0])]
    if not dashed:
        draw.line(pts, fill=rgb, width=width, joint="curve")
        return
    on, acc = True, 0.0
    for a, b in zip(pts[:-1], pts[1:]):
        seg = float(np.hypot(b[0] - a[0], b[1] - a[1]))
        if on:
            draw.line([a, b], fill=rgb, width=width)
        acc += seg
        if acc > 6:
            on, acc = not on, 0.0


def compare_image(d, cfg, images_dir, pred_dir):
    """Side-by-side PIL image and change counts for one Done tile."""
    color = {c["name"]: _rgb(c["color"]) for c in cfg["classes"]}
    w, h = d["width"], d["height"]
    with Image.open(images_dir / d["image"]) as im:
        base = im.convert("RGB")
    cut = d.get("conf_cut") or 0.0
    orig, confs = read_predictions(pred_dir / (Path(d["image"]).stem + ".txt"), w, h)
    model_cls = {int(o["id"][1:]): o["model_class"] for o in d["objects"] if o["source"] == "model"}

    left = base.copy(); dl = ImageDraw.Draw(left)
    for k, (p, c) in enumerate(zip(orig, confs)):
        if c >= cut - 1e-9:
            _poly(dl, p, color.get(model_cls.get(k), (255, 255, 255)), 3)

    changes = classify_changes(d, pred_dir)
    right = base.copy(); dr = ImageDraw.Draw(right)
    small = _font(22)
    for o, ch in sorted(changes, key=lambda t: t[1] != "rejected"):   # rejected underneath
        if ch == "pending" and o["source"] == "model" and o["conf"] < cut - 1e-9:
            continue
        pts = [[x * w, y * h] for x, y in o["poly"]]
        if ch == "kept":
            _poly(dr, pts, color[o["cls"]], 3)
        elif ch == "reclassed":   # yellow outline + the new class name
            _poly(dr, pts, CHANGE_RGB["reclassed"], 3)
            x0, y0 = min(p[0] for p in pts), min(p[1] for p in pts)
            dr.text((x0, max(0, y0 - 26)), "-> " + o["cls"], fill=CHANGE_RGB["reclassed"], font=small,
                    stroke_width=3, stroke_fill=(0, 0, 0))
        else:
            _poly(dr, pts, CHANGE_RGB[ch], 2 if ch == "pending" else 3, dashed=ch == "rejected")

    n = Counter(ch for _, ch in changes)
    title, leg = _font(26), _font(18)
    gap, top, bot = 12, 44, 40
    out = Image.new("RGB", (2 * w + gap, h + top + bot), (13, 17, 23))
    out.paste(left, (0, top)); out.paste(right, (w + gap, top))
    dd = ImageDraw.Draw(out)
    dd.text((10, 8), f"Original (model, conf >= {cut:.2f})", fill=(230, 237, 243), font=title)
    dd.text((w + gap + 10, 8), "Revised", fill=(230, 237, 243), font=title)
    dd.text((10, h + top + 10), d["image"], fill=(139, 148, 158), font=leg)
    x = w + gap + 10
    items = [("kept (class colour)", (63, 185, 80), n["kept"]), ("reclassed", CHANGE_RGB["reclassed"], n["reclassed"]),
             ("edited", CHANGE_RGB["edited"], n["edited"]), ("drawn", CHANGE_RGB["drawn"], n["drawn"]),
             ("rejected", CHANGE_RGB["rejected"], n["rejected"]), ("unreviewed", CHANGE_RGB["pending"], n["pending"])]
    for label, rgb, cnt in items:
        dd.rectangle([x, h + top + 14, x + 14, h + top + 28], fill=rgb)
        txt = f"{label} {cnt}"
        dd.text((x + 20, h + top + 10), txt, fill=(230, 237, 243), font=leg)
        x += 34 + int(dd.textlength(txt, font=leg))
    return out, n


def export_compare(docs, cfg, images_dir, pred_dir, out_dir):
    """compare/<tile>.png + compare_all.pdf. Returns [(name, png path, counts)]."""
    cdir = out_dir / "compare"
    cdir.mkdir(parents=True, exist_ok=True)
    rows, pages = [], []
    for d in docs:
        im, n = compare_image(d, cfg, images_dir, pred_dir)
        p = cdir / (Path(d["image"]).stem + ".png")
        im.save(p)
        rows.append((d["image"], p, n))
        pages.append(im)
    if pages:
        Image.init()   # PDF pages are JPEG-encoded; make sure that writer is registered
        pages[0].save(out_dir / "compare_all.pdf", save_all=True, append_images=pages[1:], resolution=150)
    return rows


# ---------------------------------------------------------------- figures

def _style(ax):
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(alpha=0.25)


def fig_threshold(by_t, analysis_conf, floor, path):
    t = [r["conf_cut"] for r in by_t]
    fig, ax = plt.subplots(figsize=(6.5, 4))
    for key, c in (("precision", "#1f77b4"), ("recall", "#d62728"), ("f1", "#2ca02c")):
        ax.plot(t, [r[key] for r in by_t], marker="o", ms=4, label=key.upper() if key == "f1" else key.title(), color=c)
    ax.axvline(analysis_conf, ls="--", color="k", lw=1)
    ax.text(analysis_conf, 0.02, f" analysis cut {analysis_conf}", fontsize=8)
    if floor > 0:
        ax.axvspan(0, floor, color="0.85", zorder=0)
        ax.text(floor / 2, 0.5, "not\nreviewed", ha="center", fontsize=8, color="0.4")
    ax.set(xlim=(0, 1), ylim=(0, 1.05), xlabel="confidence cut", ylabel="score",
           title="Model vs reviewer, by confidence cut")
    _style(ax); ax.legend(frameon=False, loc="lower left")
    fig.tight_layout(); fig.savefig(path, dpi=150); plt.close(fig)


def fig_bins(by_conf, path):
    fig, ax = plt.subplots(figsize=(6.5, 4))
    x = np.arange(len(by_conf))
    ax.bar(x, [r["precision"] for r in by_conf], color="#1f77b4")
    for i, r in enumerate(by_conf):
        ax.text(i, r["precision"] + 0.02, f"n={r['n_reviewed']}", ha="center", fontsize=8)
    ax.set_xticks(x, [r["conf_bin"] for r in by_conf], rotation=45)
    ax.set(ylim=(0, 1.12), xlabel="model confidence", ylabel="share kept as EV",
           title="Precision by confidence bin")
    _style(ax); fig.tight_layout(); fig.savefig(path, dpi=150); plt.close(fig)


def confusion_matrix(conf_rows):
    rows = sorted({r["model_class"] for r in conf_rows})
    cols = sorted({r["reviewer"] for r in conf_rows} - {"rejected"}) + ["rejected"]
    m = np.zeros((len(rows), len(cols)), int)
    for r in conf_rows:
        m[rows.index(r["model_class"]), cols.index(r["reviewer"])] += r["n"]
    return rows, cols, m


def fig_confusion(conf_rows, path):
    rows, cols, m = confusion_matrix(conf_rows)
    fig, ax = plt.subplots(figsize=(1.3 * len(cols) + 2.5, 0.7 * len(rows) + 2))
    ax.imshow(m, cmap="Blues")
    for i in range(len(rows)):
        for j in range(len(cols)):
            ax.text(j, i, m[i, j], ha="center", va="center",
                    color="white" if m[i, j] > m.max() / 2 else "black")
    ax.set_xticks(range(len(cols)), cols, rotation=30, ha="right")
    ax.set_yticks(range(len(rows)), rows)
    ax.set(xlabel="reviewer", ylabel="model", title="Class: model -> reviewer")
    fig.tight_layout(); fig.savefig(path, dpi=150); plt.close(fig)


def fig_sizes(reviewed, model, path):
    from visualization.compare_size_profiles import plot_violin
    series = {}
    for smp in sorted({r["sample"] for r in reviewed + model}):
        for tag, rows in (("model", model), ("reviewed", reviewed)):
            v = [float(r["equivalent_diameter"]) for r in rows
                 if r["sample"] == smp and r["role"] != "inner-layer" and r.get("equivalent_diameter") is not None]
            if v:
                series[f"{smp}\n{tag}"] = v
    if not series:
        return False
    fig, ax = plt.subplots(figsize=(max(5, 1.1 * len(series) + 1.5), 4.2))
    plot_violin(ax, series, "", group_label="sample")
    ax.set(ylabel="equivalent diameter (nm)", title="EV size, model vs reviewed (same tiles)")
    _style(ax); fig.tight_layout(); fig.savefig(path, dpi=150); plt.close(fig)
    return True


def make_figures(fig_dir, by_t, by_conf, conf_rows, reviewed, model, analysis_conf, floor):
    fig_dir.mkdir(parents=True, exist_ok=True)
    out = {}
    if by_t:
        fig_threshold(by_t, analysis_conf, floor, fig_dir / "pr_vs_confidence.png"); out["pr"] = fig_dir / "pr_vs_confidence.png"
    if by_conf:
        fig_bins(by_conf, fig_dir / "precision_by_conf.png"); out["bins"] = fig_dir / "precision_by_conf.png"
    if conf_rows:
        fig_confusion(conf_rows, fig_dir / "confusion.png"); out["conf"] = fig_dir / "confusion.png"
    if fig_sizes(reviewed, model, fig_dir / "sizes_model_vs_reviewed.png"):
        out["sizes"] = fig_dir / "sizes_model_vs_reviewed.png"
    return out


# ---------------------------------------------------------------- tables

def write_xlsx(path, sheets):
    """sheets: [(title, header, rows)]. Returns the path written (R8: renamed if locked)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(); wb.remove(wb.active)
    for title, header, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(header)
        for r in rows:
            ws.append(list(r))
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.freeze_panes = "A2"
        for j, _ in enumerate(header, start=1):
            vals = [ws.cell(i, j).value for i in range(1, ws.max_row + 1)]
            ws.column_dimensions[get_column_letter(j)].width = min(45, max(9, *(len(str(v)) + 2 for v in vals)))
            for i in range(2, ws.max_row + 1):
                v = ws.cell(i, j).value
                if isinstance(v, float):
                    ws.cell(i, j).number_format = "0.000"
    try:
        wb.save(path)
        return path
    except PermissionError:
        alt = path.with_name(f"{path.stem}_{time.strftime('%H%M%S')}{path.suffix}")
        wb.save(alt)
        print(f"  {path.name} is open elsewhere (Excel?) -> wrote {alt.name}")
        return alt


def _b64(path, max_w=None, fmt="PNG"):
    im = Image.open(path)
    if max_w and im.width > max_w:
        im = im.resize((max_w, round(im.height * max_w / im.width)))
    buf = io.BytesIO()
    im.convert("RGB").save(buf, fmt, quality=82) if fmt == "JPEG" else im.save(buf, fmt)
    return f"data:image/{fmt.lower()};base64," + base64.b64encode(buf.getvalue()).decode()


def _table(header, rows):
    def cell(v):
        return f"{v:.3f}" if isinstance(v, float) else html.escape(str(v))
    return ("<table><tr>" + "".join(f"<th>{html.escape(h)}</th>" for h in header) + "</tr>"
            + "".join("<tr>" + "".join(f"<td>{cell(v)}</td>" for v in r) + "</tr>" for r in rows)
            + "</table>")


def write_html(path, headline, sheets, figs, compare_rows):
    """Self-contained page (R7): figures and thumbnails embedded."""
    fig_html = "".join(f'<figure><img src="{_b64(p)}"></figure>' for p in figs.values())
    tables = "".join(f"<h3>{html.escape(t)}</h3>{_table(h, r)}" for t, h, r in sheets if r)
    thumbs = "".join(
        f'<figure><a href="../export/compare/{p.name}">'   # page lives in metrics/
        f'<img src="{_b64(p, 1100, "JPEG")}"></a><figcaption>{html.escape(name)} — '
        + ", ".join(f"{k} {v}" for k, v in sorted(n.items()) if v) + "</figcaption></figure>"
        for name, p, n in compare_rows)
    head = "".join(f"<div class='kpi'><b>{html.escape(str(v))}</b><span>{html.escape(k)}</span></div>"
                   for k, v in headline)
    page = f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>EV QC report</title>
<style>
:root{{--bg:#fff;--fg:#1f2328;--mut:#59636e;--line:#d1d9e0;--th:#f6f8fa}}
@media (prefers-color-scheme:dark){{:root{{--bg:#0d1117;--fg:#e6edf3;--mut:#8b949e;--line:#30363d;--th:#161b22}}}}
body{{background:var(--bg);color:var(--fg);font:14px system-ui,Segoe UI,sans-serif;margin:0 auto;max-width:1200px;padding:16px}}
h1{{margin:0 0 4px}} .sub{{color:var(--mut);margin-bottom:16px}}
.kpis{{display:flex;flex-wrap:wrap;gap:10px;margin:12px 0 20px}}
.kpi{{border:1px solid var(--line);border-radius:8px;padding:10px 14px;min-width:120px}}
.kpi b{{display:block;font-size:22px}} .kpi span{{color:var(--mut);font-size:12px}}
.grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,520px),1fr));gap:12px}}
figure{{margin:0 0 14px}} img{{max-width:100%;height:auto;border-radius:6px;background:#fff}}
figcaption{{color:var(--mut);font-size:12px;margin-top:4px}}
table{{border-collapse:collapse;margin:6px 0 18px;font-size:13px;display:block;overflow-x:auto}}
th,td{{border:1px solid var(--line);padding:4px 8px;text-align:right}} th{{background:var(--th)}}
td:first-child,th:first-child{{text-align:left}}
</style></head><body>
<h1>EV QC report</h1><div class="sub">{time.strftime('%Y-%m-%d %H:%M')} · images marked Done only · outlines left pending are not scored</div>
<div class="kpis">{head}</div>
<h2>Figures</h2><div class="grid">{fig_html}</div>
<h2>Tables</h2>{tables}
<h2>Original vs revised</h2>
<div class="sub">Revised: class colour = kept · yellow + label = reclassed · orange = edited · cyan = drawn · red dashed = rejected · grey = unreviewed. Click for full size.</div>
{thumbs}
</body></html>"""
    path.write_text(page, encoding="utf-8")
